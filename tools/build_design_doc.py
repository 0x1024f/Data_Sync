from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "enterprise-data-sync-architecture.md"
OUT = ROOT / "docs" / "enterprise-data-sync-architecture.docx"

BLUE = "1F4E79"
DARK = "17365D"
MUTED = "667085"
LIGHT = "EAF1F8"
FONT_CN = "STSong"


def set_font(run, size=None, bold=None, color=None, name=FONT_CN):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{slot}"), name)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal._element.rPr.rFonts.set(qn(f"w:{slot}"), FONT_CN)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (
        ("Title", 25, DARK, 0, 10),
        ("Subtitle", 11, MUTED, 0, 14),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, DARK, 8, 4),
    ):
        st = doc.styles[name]
        st.font.name = FONT_CN
        for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
            st._element.rPr.rFonts.set(qn(f"w:{slot}"), FONT_CN)
        st.font.size = Pt(size)
        st.font.bold = name != "Subtitle"
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        st = doc.styles[name]
        st.font.name = FONT_CN
        for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
            st._element.rPr.rFonts.set(qn(f"w:{slot}"), FONT_CN)
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.1


def add_inline(p, text):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            set_font(r, size=9.3, color=DARK, name="Menlo")
        elif part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            set_font(r, bold=True)
        else:
            set_font(p.add_run(part))


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    shade(cell, "F4F6F8")
    set_cell_margins(cell, 130, 150, 130, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("\n".join(lines))
    set_font(r, size=8.2, color="273142", name="Menlo")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def header_footer(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)
    hp = sec.header.paragraphs[0]
    hp.text = "DATA SYNC  ·  企业级架构设计"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(hp.runs[0], size=8.5, color=MUTED, bold=True)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("内部技术设计  |  ")
    set_font(r, size=8, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)


def build():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_styles(doc)
    header_footer(doc)
    props = doc.core_properties
    props.title = "企业级内网数据可靠同步平台设计"
    props.subject = "文件与数据库增量经前置机同步至 MinIO/目标数据库"
    props.author = "Data Sync Architecture"

    in_code = False
    code = []
    first_title = True
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                add_code(doc, code)
                code = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code.append(line)
            continue
        if not line:
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.space_after = Pt(10)
            add_inline(p, line[2:])
            for r in p.runs:
                set_font(r, size=9.3, color=MUTED)
            continue
        m = re.match(r"^(#{1,3})\s+(.+)$", line)
        if m:
            level, txt = len(m.group(1)), m.group(2)
            if level == 1 and first_title:
                p = doc.add_paragraph(style="Title")
                add_inline(p, txt)
                sub = doc.add_paragraph(style="Subtitle")
                add_inline(sub, "源端 Agent · 单端口可靠传输 · 前置机持久中继 · MinIO / 数据库落地")
                first_title = False
            else:
                p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
                add_inline(p, txt)
            continue
        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", line))
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            continue
        p = doc.add_paragraph()
        add_inline(p, line)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
